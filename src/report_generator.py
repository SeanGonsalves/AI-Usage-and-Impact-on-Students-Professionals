"""Write concise Excel, Markdown, and PDF analysis reports."""

from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd

from src.config import REPORT_DIRS


def _markdown_table(frame: pd.DataFrame) -> str:
    if frame.empty:
        return "_No results available._"
    displayed = frame.copy()
    headers = list(displayed.columns)
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for row in displayed.itertuples(index=False, name=None):
        values = []
        for value in row:
            if pd.isna(value):
                values.append("")
            elif isinstance(value, float):
                values.append(f"{value:.3g}")
            else:
                values.append(str(value).replace("|", "\\|"))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def generate_reports(results: dict, plots: list[Path], input_name: str) -> dict[str, Path]:
    """Write all reports beneath their configured format folders."""
    markdown_path = _write_markdown(results, plots, input_name)
    excel_path = _write_excel(results)
    pdf_path = _write_pdf(results, plots)
    return {"markdown": markdown_path, "excel": excel_path, "pdf": pdf_path}


def _write_markdown(results: dict, plots: list[Path], input_name: str) -> Path:
    destination = REPORT_DIRS["markdown"] / "analysis_report.md"
    image_by_name = {plot.name: plot for plot in plots}
    preferred_plots = (
        "ai_usage_hours_per_day_distribution.png",
        "ai_purpose_frequency.png",
        "key_measures_by_user_type.png",
        "ai_usage_and_time_saved.png",
    )
    images = []
    for name in preferred_plots:
        plot = image_by_name.get(name)
        if plot is not None:
            relative = Path("../../plots") / plot.parent.name / plot.name
            images.append(f"![{name.replace('_', ' ').removesuffix('.png')}]({relative.as_posix()})")

    text = [
        "# AI Usage and Impact on Students and Professionals",
        "",
        "A descriptive analysis of reported AI use and perceptions. This observational dataset "
        "cannot show that AI use caused any reported outcome.",
        "",
        "## Dataset and data quality",
        f"Source: `{input_name}`",
        "",
        _markdown_table(results["summary"]),
        "",
        "Exact duplicate rows were excluded. Text labels and supported numeric formats were "
        "standardized. Missing values were not imputed; out-of-range values were flagged "
        "and left missing.",
        "",
        _markdown_table(results["quality"]),
        "",
        "## Key findings",
        *[f"- {finding}" for finding in results["findings"]],
        "",
        "## Descriptive results",
        _markdown_table(results["descriptive"]),
        "",
        "### Common responses",
        _markdown_table(
            results["categories"].loc[
                results["categories"]["variable"].isin(
                    ["AI_Purpose", "AI_Tool", "Would_Recommend"]
                )
            ].groupby("variable", sort=False).head(8)
        ),
        "",
        "## Student and professional comparison",
        "Outcome comparisons use the Mann–Whitney U test on available responses. Adjusted "
        "p-values use the Benjamini–Hochberg method. Results are descriptive, not causal.",
        "",
        _markdown_table(results["comparisons"]),
        "",
        "AI purpose by group (percent within each user type):",
        "",
        _markdown_table(results["purpose_comparison"]),
        "",
        "## Association",
        "Spearman correlation describes a monotonic association and does not establish cause.",
        "",
        _markdown_table(results["correlation"]),
        "",
        "## Selected plots",
        "",
        *images,
        "",
        "## Limitations",
        "The CSV's source, sampling method, and license are undocumented. Responses are "
        "self-reported, missingness may affect results, and findings should not be generalized "
        "beyond these records.",
        "",
        "## Reproduce",
        "From the project root, install `requirements.txt` and run `python main.py`. "
        "Use `python main.py --input path/to/file.csv` for another CSV.",
        "",
    ]
    destination.write_text("\n".join(text), encoding="utf-8")
    return destination


def _write_excel(results: dict) -> Path:
    destination = REPORT_DIRS["excel"] / "analysis_report.xlsx"
    sheets = {
        "Summary": results["summary"],
        "Data quality": results["quality"],
        "Descriptive stats": results["descriptive"],
        "Categories": results["categories"],
        "Group comparison": results["comparisons"],
        "Purpose by group": results["purpose_comparison"],
        "Association": results["correlation"],
        "Key findings": pd.DataFrame({"finding": results["findings"]}),
    }
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            if not frame.empty:
                frame.to_excel(writer, sheet_name=name[:31], index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            for cells in worksheet.columns:
                width = min(
                    42,
                    max(10, max(len(str(cell.value or "")) for cell in cells) + 2),
                )
                worksheet.column_dimensions[get_column_letter(cells[0].column)].width = width
    return destination


def _write_pdf(results: dict, plots: list[Path]) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    destination = REPORT_DIRS["pdf"] / "analysis_report.pdf"
    document = SimpleDocTemplate(
        str(destination),
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title="AI Usage and Impact Analysis",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("Cell", parent=styles["BodyText"], fontSize=7, leading=8))
    story = [
        Paragraph("AI Usage and Impact on Students and Professionals", styles["Title"]),
        Paragraph("Research summary", styles["Heading1"]),
    ]
    for finding in results["findings"]:
        story.extend([Paragraph(escape(finding), styles["BodyText"]), Spacer(1, 4)])

    def add_table(title: str, frame: pd.DataFrame, maximum: int = 12) -> None:
        if frame.empty:
            return
        story.append(Paragraph(escape(title), styles["Heading2"]))
        frame = frame.head(maximum)
        data = [[Paragraph(escape(str(value)), styles["Cell"]) for value in frame.columns]]
        for row in frame.itertuples(index=False, name=None):
            data.append([
                Paragraph(
                    escape("" if pd.isna(value) else (
                        f"{value:.3g}" if isinstance(value, float) else str(value)
                    )),
                    styles["Cell"],
                )
                for value in row
            ])
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CCD5E0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
        ]))
        story.extend([table, Spacer(1, 8)])

    add_table("Dataset", results["summary"])
    add_table("Descriptive statistics", results["descriptive"])
    add_table("Student and professional comparison", results["comparisons"])
    add_table("AI purpose by group (%)", results["purpose_comparison"])
    add_table("AI use and time saved", results["correlation"])
    story.append(Paragraph("Selected visualizations", styles["Heading1"]))
    image_by_name = {plot.name: plot for plot in plots}
    for name in (
        "ai_usage_hours_per_day_distribution.png",
        "ai_purpose_frequency.png",
        "key_measures_by_user_type.png",
        "ai_usage_and_time_saved.png",
    ):
        plot = image_by_name.get(name)
        if plot and plot.is_file():
            story.extend([
                Image(str(plot), width=165 * mm, height=85 * mm, kind="proportional"),
                Spacer(1, 5),
            ])
    story.extend([
        Paragraph("Limitations", styles["Heading1"]),
        Paragraph(
            "The sampling method and dataset provenance are unknown. Responses are self-reported, "
            "missing data was not imputed, and associations do not demonstrate causation.",
            styles["BodyText"],
        ),
    ])
    document.build(story)
    return destination
